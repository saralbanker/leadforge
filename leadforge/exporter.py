import pandas as pd
from pathlib import Path
from typing import List, Dict, Any
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from leadforge.config import OUTPUT_DIR
from leadforge.utils import get_logger

logger = get_logger()

def export_leads_to_excel(leads: List[Dict[str, Any]], filename: str) -> Path:
    """
    Exports the scored leads into a styled, professional Excel (.xlsx) file.
    """
    if not filename.endswith(".xlsx"):
        filename += ".xlsx"

    output_path = OUTPUT_DIR / filename

    # 1. Convert to DataFrame
    df = pd.DataFrame(leads)

    # Default columns mapping
    column_mapping = {
        "name": "Business Name",
        "category": "Business Category",
        "phone": "Phone Number",
        "website": "Website",
        "address": "Address",
        "area": "Area",
        "priority": "Priority",
        "notes": "Notes",
        "discovery_date": "Discovery Date"
    }

    # Ensure all target columns exist, fill missing with empty
    for col in column_mapping.keys():
        if col not in df.columns:
            df[col] = ""

    # Filter and reorder columns
    df_filtered = df[list(column_mapping.keys())].rename(columns=column_mapping)

    # 2. Write to Excel
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_filtered.to_excel(writer, index=False, sheet_name="Leads")

        # Access openpyxl sheet to apply professional styling
        worksheet = writer.sheets["Leads"]

        # Color palettes
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # Navy
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        high_priority_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # Soft orange/red
        medium_priority_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Soft green

        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        # Format headers
        for col_idx in range(1, len(column_mapping) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # Format data cells
        for row_idx in range(2, len(leads) + 2):
            priority_val = worksheet.cell(row=row_idx, column=7).value  # Column 7 is "Priority"

            # Determine fill color for the row based on priority
            row_fill = None
            if priority_val == "High":
                row_fill = high_priority_fill
            elif priority_val == "Medium":
                row_fill = medium_priority_fill

            for col_idx in range(1, len(column_mapping) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = Font(name="Calibri", size=11)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

                # Apply priority coloring to the Priority cell specifically, or subtle row highlight
                if col_idx == 7 and row_fill:
                    cell.fill = row_fill
                    cell.font = Font(name="Calibri", size=11, bold=True, color="333333")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx in [3, 9]:  # Center align Phone Number and Discovery Date
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        # Set row heights
        worksheet.row_dimensions[1].height = 28
        for r in range(2, len(leads) + 2):
            worksheet.row_dimensions[r].height = 20

        # Adjust column widths dynamically
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)

            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)

            # Add padding and limit width
            worksheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    logger.info(f"Successfully exported {len(leads)} leads to: {output_path}")
    return output_path


def export_intelligence_to_excel(leads: List[Dict[str, Any]], filename: str) -> Path:
    """Export enriched intelligence data with all MVP columns."""
    if not filename.endswith(".xlsx"):
        filename += ".xlsx"
    output_path = OUTPUT_DIR / filename

    column_mapping = {
        "name":                  "Business Name",
        "phone":                 "Phone",
        "website":               "Website",
        "category":              "Category",
        "area":                  "Area",
        "rating":                "Rating",
        "review_count":          "Review Count",
        "maturity_grade":        "Digital Maturity",
        "priority":              "Priority",
        "score":                 "Opportunity Score",
        "top_opportunity":       "Top Opportunity",
        "recommended_services":  "Recommended Services",
    }

    df = pd.DataFrame(leads)
    for col in column_mapping:
        if col not in df.columns:
            df[col] = ""
    df_out = df[list(column_mapping)].rename(columns=column_mapping)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_out.to_excel(writer, index=False, sheet_name="Intelligence Export")
        ws = writer.sheets["Intelligence Export"]

        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        high_fill   = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        medium_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        for col_idx in range(1, len(column_mapping) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        priority_col = list(column_mapping).index("priority") + 1

        for row_idx in range(2, len(leads) + 2):
            priority_val = ws.cell(row=row_idx, column=priority_col).value
            row_fill = high_fill if priority_val == "High" else (medium_fill if priority_val == "Medium" else None)

            for col_idx in range(1, len(column_mapping) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = Font(name="Calibri", size=10)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
                if col_idx == priority_col and row_fill:
                    cell.fill = row_fill
                    cell.font = Font(name="Calibri", size=10, bold=True)
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[1].height = 28
        for r in range(2, len(leads) + 2):
            ws.row_dimensions[r].height = 18

        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    logger.info(f"Intelligence export: {len(leads)} rows → {output_path}")
    return output_path
